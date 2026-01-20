import unittest
import os
from office_lib import OfficeAutomator
from docx import Document
from docx.oxml.ns import qn

class TestOfficeAutomator(unittest.TestCase):
    def setUp(self):
        self.test_file = "test_doc_sec.docx"
        self.doc = Document()

        # Section 1
        self.doc.add_paragraph("Section 1 Content")

        # Section 2
        self.doc.add_section()
        self.doc.add_paragraph("Section 2 Content")

        # Bold text for formatting test
        p = self.doc.add_paragraph()
        run = p.add_run("  Bold  Text  ")
        run.bold = True

        self.doc.save(self.test_file)

        self.automator = OfficeAutomator()
        self.automator.load_document(self.test_file)

    def tearDown(self):
        if os.path.exists(self.test_file):
            os.remove(self.test_file)

    def test_page_number_continuity(self):
        # Start at 10 on first section
        self.automator.add_page_number(start_at=10)

        # Check Section 1
        sect1 = self.automator.doc.sections[0]
        pgNumType1 = sect1._sectPr.find(qn('w:pgNumType'))
        self.assertIsNotNone(pgNumType1)
        self.assertEqual(pgNumType1.get(qn('w:start')), "10")

        # Check Section 2 (Should NOT have w:start set)
        sect2 = self.automator.doc.sections[1]
        pgNumType2 = sect2._sectPr.find(qn('w:pgNumType'))
        if pgNumType2:
            self.assertIsNone(pgNumType2.get(qn('w:start')))

    def test_clean_spaces_preserve_bold(self):
        # Find the bold paragraph by content to be sure
        target_p = None
        target_run = None
        for p in self.automator.doc.paragraphs:
            for r in p.runs:
                if "Bold" in r.text:
                    target_p = p
                    target_run = r
                    break
            if target_run: break

        self.assertIsNotNone(target_run, "Could not find bold run")
        # Verify initial state
        # In some python-docx versions/XML states, boolean True is returned.
        # If None, it means inherited. But we saved it as True.
        self.assertTrue(target_run.bold, f"Initial bold state is {target_run.bold}")

        self.automator.clean_extra_spaces()

        # Verify text cleaned
        self.assertIn("Bold Text", target_run.text)
        # Verify formatting preserved
        self.assertTrue(target_run.bold, f"Post-clean bold state is {target_run.bold}")

if __name__ == '__main__':
    unittest.main()
