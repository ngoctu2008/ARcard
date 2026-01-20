import os
import re
from docx import Document
from docx.shared import Pt, Cm, Mm
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import platform

# Try importing docx2pdf, handle failure gracefully if not installed/supported
try:
    from docx2pdf import convert as convert_pdf
    HAS_PDF_SUPPORT = True
except ImportError:
    HAS_PDF_SUPPORT = False

class OfficeAutomator:
    def __init__(self):
        self.doc = None
        self.doc_path = None

    def load_document(self, path):
        """Loads a .docx document from the given path."""
        if not os.path.exists(path):
            raise FileNotFoundError(f"File not found: {path}")

        try:
            self.doc = Document(path)
            self.doc_path = path
            return True
        except Exception as e:
            raise Exception(f"Error loading document: {e}")

    def save_document(self, path=None):
        """Saves the document to the given path or overwrites the original."""
        if self.doc is None:
            raise Exception("No document loaded.")

        target_path = path if path else self.doc_path
        try:
            self.doc.save(target_path)
            return True
        except Exception as e:
            raise Exception(f"Error saving document: {e}")

    # --- Formatting ---

    def set_font_style(self, font_name="Times New Roman", font_size=14):
        if self.doc is None: raise Exception("No document loaded.")

        def apply_font(run):
            if font_name:
                run.font.name = font_name
                rPr = run._element.get_or_add_rPr()
                rFonts = rPr.get_or_add_rFonts()
                rFonts.set(qn('w:eastAsia'), font_name)
            if font_size:
                try:
                    run.font.size = Pt(float(font_size))
                except ValueError:
                    pass

        for paragraph in self.doc.paragraphs:
            for run in paragraph.runs:
                apply_font(run)

        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            apply_font(run)
        return True

    def set_paragraph_style(self, line_spacing=1.5, space_before=0, space_after=6):
        if self.doc is None: raise Exception("No document loaded.")

        def apply_para(p):
            if line_spacing:
                try:
                    p.paragraph_format.line_spacing = float(line_spacing)
                except ValueError:
                    pass
            if space_before is not None:
                try:
                    p.paragraph_format.space_before = Pt(float(space_before))
                except ValueError:
                    pass
            if space_after is not None:
                try:
                    p.paragraph_format.space_after = Pt(float(space_after))
                except ValueError:
                    pass

        for paragraph in self.doc.paragraphs:
            apply_para(paragraph)

        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        apply_para(paragraph)
        return True

    def set_alignment(self, align_type):
        """
        Sets paragraph alignment.
        align_type: 'LEFT', 'CENTER', 'RIGHT', 'JUSTIFY'
        """
        if self.doc is None: raise Exception("No document loaded.")

        align_map = {
            'LEFT': WD_ALIGN_PARAGRAPH.LEFT,
            'CENTER': WD_ALIGN_PARAGRAPH.CENTER,
            'RIGHT': WD_ALIGN_PARAGRAPH.RIGHT,
            'JUSTIFY': WD_ALIGN_PARAGRAPH.JUSTIFY
        }

        wd_align = align_map.get(align_type.upper())
        if wd_align is None: return False

        for paragraph in self.doc.paragraphs:
            paragraph.alignment = wd_align

        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.alignment = wd_align
        return True

    def set_indentation(self, indent_type, indent_value_cm):
        if self.doc is None: raise Exception("No document loaded.")

        val = Cm(float(indent_value_cm)) if indent_value_cm is not None else Cm(0)

        def apply_indent(p):
            if indent_type == 'None' or indent_type == '(Không)':
                p.paragraph_format.first_line_indent = 0
                p.paragraph_format.left_indent = 0
            elif indent_type == 'First Line':
                p.paragraph_format.first_line_indent = val
                p.paragraph_format.left_indent = 0
            elif indent_type == 'Hanging':
                p.paragraph_format.first_line_indent = -val
                p.paragraph_format.left_indent = val

        for paragraph in self.doc.paragraphs:
            apply_indent(paragraph)

        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        apply_indent(paragraph)
        return True

    # --- Page Setup ---

    def set_orientation_whole_doc(self, orientation):
        if self.doc is None: raise Exception("No document loaded.")

        for section in self.doc.sections:
            current_orient = section.orientation
            current_width = section.page_width
            current_height = section.page_height

            if orientation == 'LANDSCAPE':
                if current_orient == WD_ORIENT.PORTRAIT or current_orient is None:
                    section.orientation = WD_ORIENT.LANDSCAPE
                    section.page_width = current_height
                    section.page_height = current_width
            elif orientation == 'PORTRAIT':
                if current_orient == WD_ORIENT.LANDSCAPE:
                    section.orientation = WD_ORIENT.PORTRAIT
                    section.page_width = current_height
                    section.page_height = current_width
        return True

    def set_paper_size(self, size_name):
        if self.doc is None: raise Exception("No document loaded.")

        sizes = {
            'A4': (210, 297),
            'Letter': (215.9, 279.4),
            'Legal': (215.9, 355.6),
            'A3': (297, 420),
            'A5': (148, 210)
        }

        key = None
        for k in sizes:
            if k.lower() == size_name.lower():
                key = k
                break

        if not key: return False

        width_mm, height_mm = sizes[key]

        for section in self.doc.sections:
            if section.orientation == WD_ORIENT.LANDSCAPE:
                section.page_width = Mm(height_mm)
                section.page_height = Mm(width_mm)
            else:
                section.page_width = Mm(width_mm)
                section.page_height = Mm(height_mm)
        return True

    def set_margins(self, top, bottom, left, right):
        if self.doc is None: raise Exception("No document loaded.")

        for section in self.doc.sections:
            if top: section.top_margin = Cm(float(top))
            if bottom: section.bottom_margin = Cm(float(bottom))
            if left: section.left_margin = Cm(float(left))
            if right: section.right_margin = Cm(float(right))
        return True

    # --- Page Numbering ---

    def add_page_number(self, align='CENTER', start_at=None, skip_first=False):
        if self.doc is None: raise Exception("No document loaded.")

        alignment_map = {
            'LEFT': WD_ALIGN_PARAGRAPH.LEFT,
            'CENTER': WD_ALIGN_PARAGRAPH.CENTER,
            'RIGHT': WD_ALIGN_PARAGRAPH.RIGHT
        }
        wd_align = alignment_map.get(str(align).upper(), WD_ALIGN_PARAGRAPH.CENTER)

        for i, section in enumerate(self.doc.sections):
            # Start At Logic
            # Only apply start_at to the first section to avoid forcing restart on every section
            if start_at is not None and i == 0:
                sectPr = section._sectPr
                pgNumType = sectPr.find(qn('w:pgNumType'))
                if pgNumType is None:
                    pgNumType = OxmlElement('w:pgNumType')
                    sectPr.append(pgNumType)
                pgNumType.set(qn('w:start'), str(int(start_at)))

            # Skip First Page Logic (Different First Page)
            if skip_first and i == 0:
                section.different_first_page_header_footer = True
                # No footer on first page automatically if we don't add it to first_page_footer
            else:
                section.different_first_page_header_footer = False

            footer = section.footer
            if not footer.paragraphs:
                p = footer.add_paragraph()
            else:
                p = footer.paragraphs[0]

            p.alignment = wd_align
            p.clear()

            run = p.add_run()

            fldChar1 = OxmlElement('w:fldChar')
            fldChar1.set(qn('w:fldCharType'), 'begin')

            instrText = OxmlElement('w:instrText')
            instrText.set(qn('xml:space'), 'preserve')
            instrText.text = "PAGE"

            fldChar2 = OxmlElement('w:fldChar')
            fldChar2.set(qn('w:fldCharType'), 'end')

            run._r.append(fldChar1)
            run._r.append(instrText)
            run._r.append(fldChar2)

        return True

    # --- Cleaning & Utilities ---

    def clean_extra_spaces(self):
        if self.doc is None: raise Exception("No document loaded.")

        # We iterate over runs to preserve formatting (Bold, Italic, Color)
        # Replacing paragraph.text wipes out all run formatting!

        for paragraph in self.doc.paragraphs:
            for run in paragraph.runs:
                if run.text:
                    # Clean internal spaces within the run
                    run.text = re.sub(r'\s+', ' ', run.text)

        # Tables
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            if run.text:
                                run.text = re.sub(r'\s+', ' ', run.text)
        return True

    def clean_empty_lines(self):
        if self.doc is None: raise Exception("No document loaded.")

        # Iterating backwards is safer when removing
        # But python-docx doesn't easily support removing paragraphs from the list directly
        # We can clear content or delete the element

        def is_empty(p):
            # Check if text is empty and no meaningful xml (like images)
            # Simplistic check: empty text
            return not p.text.strip()

        # We collect p elements to delete
        to_delete = []
        for p in self.doc.paragraphs:
            if is_empty(p):
                to_delete.append(p)

        for p in to_delete:
            p._element.getparent().remove(p._element)
            p._p = p._element = None

        return True

    def remove_hyperlinks(self):
        if self.doc is None: raise Exception("No document loaded.")

        # Hyperlinks are usually w:hyperlink
        # We want to replace w:hyperlink with its children (runs)

        for paragraph in self.doc.paragraphs:
            self._flatten_hyperlinks(paragraph)

        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._flatten_hyperlinks(paragraph)
        return True

    def _flatten_hyperlinks(self, paragraph):
        # Access xml directly
        p = paragraph._element
        hyperlinks = p.xpath('.//w:hyperlink')
        for hyperlink in hyperlinks:
            # Move children (runs) to parent (paragraph) before the hyperlink
            parent = hyperlink.getparent()
            index = parent.index(hyperlink)
            for child in hyperlink:
                parent.insert(index, child)
                index += 1
            parent.remove(hyperlink)

    def resize_images(self, max_width_cm):
        if self.doc is None: raise Exception("No document loaded.")

        max_width = Cm(float(max_width_cm))

        for shape in self.doc.inline_shapes:
            if shape.width > max_width:
                # Maintain aspect ratio
                ratio = max_width / shape.width
                shape.width = int(shape.width * ratio)
                shape.height = int(shape.height * ratio)
        return True

    def convert_vni_to_unicode(self):
        if self.doc is None: raise Exception("No document loaded.")
        # Placeholder
        return True

    def export_to_pdf(self, output_path=None):
        if not HAS_PDF_SUPPORT:
            raise Exception("docx2pdf is not installed or supported on this system.")

        if self.doc is None: raise Exception("No document loaded.")
        if not self.doc_path: raise Exception("Document must be saved/loaded from disk first.")

        target = output_path if output_path else self.doc_path.replace(".docx", ".pdf")

        abs_input = os.path.abspath(self.doc_path)
        abs_output = os.path.abspath(target)

        try:
            convert_pdf(abs_input, abs_output)
            return abs_output
        except Exception as e:
            raise Exception(f"PDF Conversion Failed: {e}")

    # --- Security & Info ---

    def set_file_properties(self, author=None, title=None, remove_personal_info=False):
        if self.doc is None: raise Exception("No document loaded.")

        if author is not None:
            self.doc.core_properties.author = author
        if title is not None:
            self.doc.core_properties.title = title

        if remove_personal_info:
            # Add <w:removePersonalInformation/> to settings.xml
            settings = self.doc.settings.element
            # Check if exists
            rpi = settings.find(qn('w:removePersonalInformation'))
            if rpi is None:
                rpi = OxmlElement('w:removePersonalInformation')
                settings.append(rpi)

        return True

    def set_header_footer_text(self, header_text=None, footer_text=None):
        if self.doc is None: raise Exception("No document loaded.")

        for section in self.doc.sections:
            if header_text is not None:
                h = section.header
                h.is_linked_to_previous = False
                if h.paragraphs:
                    h.paragraphs[0].text = header_text
                    for p in h.paragraphs[1:]:
                        p._element.getparent().remove(p._element)
                else:
                    h.add_paragraph(header_text)

            if footer_text is not None:
                f = section.footer
                f.is_linked_to_previous = False
                if f.paragraphs:
                    f.paragraphs[0].text = footer_text
                    for p in f.paragraphs[1:]:
                        p._element.getparent().remove(p._element)
                else:
                    f.add_paragraph(footer_text)
        return True

    def protect_document(self, protection_type='NONE', password=''):
        """
        protection_type: 'NONE', 'READ_ONLY'
        """
        if self.doc is None: raise Exception("No document loaded.")

        settings = self.doc.settings.element
        doc_protect = settings.find(qn('w:documentProtection'))

        if protection_type == 'NONE':
            if doc_protect is not None:
                settings.remove(doc_protect)
        elif protection_type == 'READ_ONLY':
            if doc_protect is None:
                doc_protect = OxmlElement('w:documentProtection')
                settings.append(doc_protect)

            doc_protect.set(qn('w:edit'), 'readOnly')
            doc_protect.set(qn('w:enforcement'), '1')

        return True

    def add_table_of_contents(self):
        """
        Inserts a Table of Contents (TOC) at the start of the document.
        Note: The TOC needs to be updated by the user (Right-click -> Update Field)
        in Word, as python-docx cannot generate page numbers/content dynamically.
        """
        if self.doc is None: raise Exception("No document loaded.")

        # Add paragraph at the beginning
        p = self.doc.add_paragraph()

        # Move to top
        body = self.doc._element.body
        body.insert(0, p._element)

        run = p.add_run()

        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')

        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = 'TOC \\o "1-3" \\h \\z \\u'

        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'separate')

        fldChar3 = OxmlElement('w:fldChar')
        fldChar3.set(qn('w:fldCharType'), 'end')

        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
        run._r.append(fldChar3)

        return True

    def get_document_info(self):
        if self.doc is None:
            return "No document loaded"
        return f"Loaded: {self.doc_path}, Sections: {len(self.doc.sections)}"
