from docx import Document
from docx.shared import Pt

def create_template(filename, content_data):
    doc = Document()

    # Set default style to Times New Roman
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    doc.add_heading('Mẫu Văn Bản Demo', 0)

    for paragraph_text in content_data:
        p = doc.add_paragraph(paragraph_text)
        # Enforce font on runs just in case
        for run in p.runs:
            run.font.name = 'Times New Roman'

    doc.save(filename)
    print(f"Đã tạo {filename}")

# Template 1: Contract style
content1 = [
    "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
    "Độc lập - Tự do - Hạnh phúc",
    "-------------------",
    "HỢP ĐỒNG LAO ĐỘNG",
    "",
    "Bên A: Công ty XYZ",
    "Đại diện bởi: {{CHUC_DANH}}",
    "Địa chỉ: {{DIA_CHI}}",
    "",
    "Bên B: Ông Nguyễn Văn A",
    "Cam kết thực hiện đúng quy định."
]

# Template 2: Notice style
content2 = [
    "THÔNG BÁO",
    "V/v: Thay đổi địa điểm làm việc",
    "",
    "Kính gửi: Toàn thể nhân viên",
    "",
    "Hiện tại, văn phòng của chúng ta tại {{DIA_CHI}} đang được tu sửa.",
    "Yêu cầu ông/bà giám đốc {{CHUC_DANH}} chỉ đạo việc di dời.",
    "",
    "Trân trọng."
]

if __name__ == "__main__":
    create_template("templates/template_hop_dong.docx", content1)
    create_template("templates/template_thong_bao.docx", content2)
